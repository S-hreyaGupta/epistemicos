#!/usr/bin/env python3
"""Negative controls for the protocol conversion checks.

Each completeness check is demonstrated failing on a document built to break
exactly that check, and a clean document is demonstrated passing.

Why this file exists: running the real truncated source through the checker
fired only ONE of the four completeness checks. Its dangling "V2" carries no
pipe, so it never registered as a verification row at all. The other three were
claims about failure modes nobody had seen fail. That is the defect the v1.0
milestone found thirteen times, so they get fixtures.

    python scripts/test_convert_protocol.py

Exit 0 = every expectation held.
"""

from __future__ import annotations

import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

SRC = Path(__file__).resolve().parent
CONVERTER = SRC / "convert_protocol.py"
TERMINAL = "Preservation verdict"

try:
    from docx import Document
except ImportError:
    print("python-docx is required:  pip install python-docx --break-system-packages",
          file=sys.stderr)
    raise SystemExit(2)

HEADER = "ID| Required element| Source| Implemented in workflow| Verification criterion| Status"


def row(n: int, fields: int = 6) -> str:
    cells = [f"V{n:02d}"] + [f"cell {n}.{k}" for k in range(1, fields)]
    return "| ".join(cells)


def build(path: Path, rows: list[str], tail: list[str]) -> None:
    d = Document()
    for line in (["Protocol", "", "8. FINAL FREEZE / RELEASE", "",
                  "Verification table changes and preservation", "", HEADER] + rows + [""] + tail):
        d.add_paragraph(line)
    d.save(str(path))


def run(path: Path) -> subprocess.CompletedProcess:
    return subprocess.run(
        [sys.executable, str(CONVERTER), "--docx", str(path),
         "--out", str(path.parent / "out.md"), "--terminal", TERMINAL, "--dry-run"],
        capture_output=True, text=True)


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="protoconv-")).resolve()
    failures: list[str] = []

    def case(label: str, needle: str, rows: list[str], tail: list[str],
             expect_pass: bool = False) -> None:
        p = tmp / (label.replace(" ", "_").replace("/", "_") + ".docx")
        build(p, rows, tail)
        r = run(p)
        out = r.stdout + r.stderr
        if expect_pass:
            if r.returncode != 0:
                failures.append(f"{label}: expected PASS, got:\n{out}")
            else:
                print(f"  [ok] passes: {label}")
            return
        if r.returncode == 0:
            failures.append(f"{label}: expected FAIL, checker passed it\n{out}")
            return
        line = next((l for l in out.splitlines()
                     if "[FAIL]" in l and needle.lower() in l.lower()), None)
        if line is None:
            failures.append(f"{label}: failed, but not on the expected check\n"
                            f"  wanted a FAIL mentioning {needle!r}\n{out}")
            return
        print(f"  [ok] fails on {needle}: {label}")

    clean = [row(n) for n in range(1, 8)]
    verdict = ["Preservation verdict", "", "Original workflow elements dropped: 0"]

    case("complete document", "", clean, verdict, expect_pass=True)

    # C4 — the shape of the real defect: rows fine, document stops after them.
    case("stops after the last row", "terminal marker", clean, [])
    case("trailing fragment with no pipe", "terminal marker", clean, ["V8"])

    # C3 — a row cut in half mid-write. Contiguity cannot see this.
    case("last row cut mid-write", "fields", clean + [row(8, fields=3)], verdict)

    # C2 — a row missing from the middle.
    case("gap in the sequence", "no gaps",
         [row(n) for n in (1, 2, 3, 5, 6, 7)], verdict)

    # C1 — out of order, and not starting at V01.
    case("rows out of order", "increase from V01",
         [row(1), row(2), row(4), row(3), row(5)], verdict)
    case("table does not start at V01", "increase from V01",
         [row(n) for n in range(3, 9)], verdict)

    # Header absent entirely.
    p = tmp / "no_header.docx"
    d = Document()
    for line in ["Protocol", "", "V01| a| b| c| d| e", "", TERMINAL]:
        d.add_paragraph(line)
    d.save(str(p))
    r = run(p)
    if r.returncode == 0 or "verification table header found" not in r.stdout.lower():
        failures.append(f"missing header: expected that check to fail\n{r.stdout}")
    else:
        print("  [ok] fails on header: verification table absent")

    # Fidelity and structure are exercised directly. The converter is
    # deterministic, so their failure modes cannot be produced by feeding it a
    # different document; they have to be fed a wrong conversion.
    sys.path.insert(0, str(SRC))
    import convert_protocol as cp

    good = cp.check_fidelity(["alpha beta", "gamma"], "alpha beta\n\ngamma\n")
    bad = cp.check_fidelity(["alpha beta", "gamma"], "alpha beta\n\ndelta\n")
    if not good.ok:
        failures.append("fidelity rejected an identical word sequence")
    elif bad.ok:
        failures.append("fidelity accepted a changed word; the check cannot fail")
    else:
        print("  [ok] fidelity accepts identical text and rejects a changed word")

    # The real regression: an earlier converter promoted numbered list items to
    # headings and fidelity passed it, because '##' contributes no words. This
    # replays that output and requires the structure check to catch it.
    src = ["intro", "", "1. STOP implementation.", "2. Record:"]
    promoted = "intro\n\n## 1. STOP implementation.\n## 2. Record:\n"
    dropped = "intro\n"
    if not cp.check_structure(src, cp.to_markdown(src)).ok:
        failures.append("structure rejected its own converter's output")
    elif cp.check_fidelity(src, promoted).ok is False:
        failures.append("fixture is wrong: fidelity should pass on invented headings")
    elif cp.check_structure(src, promoted).ok:
        failures.append("structure accepted invented headings; the check cannot fail")
    elif cp.check_structure(src, dropped).ok:
        failures.append("structure accepted dropped lines; the check cannot fail")
    else:
        print("  [ok] structure catches invented headings that fidelity passes")
        print("  [ok] structure catches dropped lines")

    shutil.rmtree(tmp, ignore_errors=True)

    print()
    if failures:
        for f in failures:
            print("FAIL:", f)
        return 1
    print("all controls fired; every completeness check is falsifiable")
    return 0


if __name__ == "__main__":
    sys.exit(main())
